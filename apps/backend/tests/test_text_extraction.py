import pytest

from src.utils.text_extraction import (
    _normalize_whitespace,
    detect_file_type,
    extract_text_from_docx,
    extract_text_from_pdf,
)


class TestNormalizeWhitespace:
    def test_collapses_multiple_spaces(self):
        assert _normalize_whitespace("hello   world") == "hello world"

    def test_collapses_multiple_newlines(self):
        result = _normalize_whitespace("line1\n\n\n\nline2")
        assert result == "line1\n\nline2"

    def test_trims_lines(self):
        result = _normalize_whitespace("  hello  \n  world  ")
        assert result == "hello\nworld"

    def test_replaces_tabs(self):
        result = _normalize_whitespace("hello\tworld")
        assert result == "hello world"

    def test_empty_string(self):
        assert _normalize_whitespace("") == ""


class TestDetectFileType:
    def test_pdf_detection(self):
        assert detect_file_type(b"%PDF-1.4 rest of content") == "pdf"

    def test_docx_detection(self):
        assert detect_file_type(b"PK\x03\x04rest of zip") == "docx"

    def test_unknown_detection(self):
        assert detect_file_type(b"random bytes") == "unknown"

    def test_empty_content(self):
        assert detect_file_type(b"") == "unknown"


class TestPdfExtraction:
    def test_invalid_pdf_raises(self):
        with pytest.raises(Exception):
            extract_text_from_pdf(b"not a valid pdf")

    def test_minimal_pdf(self):
        """Create a minimal valid PDF to test extraction."""
        import fitz

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "John Doe\nSoftware Engineer\nPython, Java, Go")
        pdf_bytes = doc.tobytes()
        doc.close()

        text = extract_text_from_pdf(pdf_bytes)
        assert "John Doe" in text
        assert "Software Engineer" in text
        assert "Python" in text

    def test_multipage_pdf(self):
        """Test extraction from a multi-page PDF."""
        import fitz

        doc = fitz.open()
        page1 = doc.new_page()
        page1.insert_text((72, 72), "Page 1: Experience")
        page2 = doc.new_page()
        page2.insert_text((72, 72), "Page 2: Education")
        pdf_bytes = doc.tobytes()
        doc.close()

        text = extract_text_from_pdf(pdf_bytes)
        assert "Page 1" in text
        assert "Page 2" in text


class TestDocxExtraction:
    def test_simple_docx(self):
        """Create a simple DOCX and test extraction."""
        from docx import Document

        doc = Document()
        doc.add_heading("John Doe", level=1)
        doc.add_paragraph("Senior Software Engineer")
        doc.add_paragraph("Python, JavaScript, Go")

        from io import BytesIO

        buffer = BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()

        text = extract_text_from_docx(docx_bytes)
        assert "John Doe" in text
        assert "Senior Software Engineer" in text
        assert "Python" in text

    def test_docx_with_bullet_points(self):
        from io import BytesIO

        from docx import Document

        doc = Document()
        doc.add_heading("Skills", level=2)
        doc.add_paragraph("Python", style="List Bullet")
        doc.add_paragraph("Java", style="List Bullet")

        buffer = BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()

        text = extract_text_from_docx(docx_bytes)
        assert "Skills" in text
        assert "Python" in text
        assert "Java" in text

    def test_docx_with_table(self):
        from io import BytesIO

        from docx import Document

        doc = Document()
        doc.add_heading("Experience", level=1)
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Company"
        table.cell(0, 1).text = "Role"
        table.cell(1, 0).text = "Acme Corp"
        table.cell(1, 1).text = "Engineer"

        buffer = BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()

        text = extract_text_from_docx(docx_bytes)
        assert "Experience" in text
        assert "Acme Corp" in text
        assert "Engineer" in text


class TestInvalidInputs:
    def test_empty_pdf(self):
        with pytest.raises(Exception):
            extract_text_from_pdf(b"")

    def test_empty_docx(self):
        with pytest.raises(Exception):
            extract_text_from_docx(b"")
