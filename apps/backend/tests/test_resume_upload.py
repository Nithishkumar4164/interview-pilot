"""Resume upload through the real endpoint: POST /api/upload/resume, then poll its status.

The model is a stand-in that behaves like the real provider: it has a full
answer of a given length and, when max_tokens is lower, returns the answer
cut off at the limit with ``truncated`` set (as Anthropic does on
stop_reason=max_tokens).
"""

import asyncio
import io
import json

import httpx
import pytest
import redis

from src.config import settings
from src.llm.interface import AnalysisContext, LLMConfig, StructuredOutput

# Measured on 2026-10-08 with claude-sonnet-5 on a synthetic 5-job resume:
# a complete extraction needed 3,050 output tokens (15 verifiable claims).
MEASURED_ANSWER_TOKENS = 3050

ANSWER = {
    "name": "Priya Raman",
    "current_role": "Staff Engineer",
    "years_experience": 8,
    "companies": ["Lumen Payments", "ShipFast Logistics"],
    "tech_stack": ["Go", "Kafka", "PostgreSQL"],
    "key_projects": [],
    "education": [],
    "certifications": [],
    "verifiable_claims": [
        {"claim": "Cut p99 latency from 820ms to 190ms", "probe_question": "What was the bottleneck?"},
    ],
}


def _redis_available() -> bool:
    try:
        return bool(redis.Redis.from_url(settings.redis_url, socket_connect_timeout=1).ping())
    except redis.RedisError:
        return False


pytestmark = pytest.mark.skipif(not _redis_available(), reason="Redis not reachable")


class _ModelWithAnswerOfLength:
    """Answers resume extraction in ``tokens`` output tokens, cut off at max_tokens like a real provider."""

    def __init__(self, tokens: int) -> None:
        self.tokens = tokens

    async def analyze(self, context: AnalysisContext, config: LLMConfig) -> StructuredOutput:
        full = "```json\n" + json.dumps(ANSWER, indent=2) + "\n```"
        if config.max_tokens >= self.tokens:
            return StructuredOutput(content=full, parsed=ANSWER, output_tokens=self.tokens)
        cut = full[: len(full) * config.max_tokens // self.tokens]
        return StructuredOutput(content=cut, parsed=None, output_tokens=config.max_tokens, truncated=True)


def _resume_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Priya Raman (synthetic)", 0)
    doc.add_paragraph("Staff Engineer, Lumen Payments. Cut p99 payment latency from 820ms to 190ms.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def _upload(monkeypatch, model: _ModelWithAnswerOfLength) -> dict:
    from src.api import upload
    from src.main import app
    from src.services.resume_service import ResumeService

    monkeypatch.setattr(upload, "_get_resume_service", lambda: ResumeService(storage=None, llm=model))  # type: ignore[arg-type]
    docx_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post("/api/upload/resume", files={"file": ("resume.docx", _resume_docx(), docx_type)})
        assert response.status_code == 200, response.text
        ref = response.json()["reference_id"]
        for _ in range(100):
            status = (await client.get(f"/api/upload/{ref}/status")).json()
            if status["status"] != "processing":
                return status
            await asyncio.sleep(0.02)
    pytest.fail("extraction never finished")


async def test_a_resume_needing_the_measured_answer_length_is_extracted(monkeypatch) -> None:
    """Before the fix the limit was 2,000 tokens: the real model's answer was cut off and the upload failed."""
    model = _ModelWithAnswerOfLength(MEASURED_ANSWER_TOKENS)
    status = await _upload(monkeypatch, model)
    assert status["status"] == "completed", status
    assert status["structured_data"]["verifiable_claims"] == ANSWER["verifiable_claims"]


async def test_an_answer_cut_off_at_the_limit_fails_with_a_clear_error(monkeypatch) -> None:
    model = _ModelWithAnswerOfLength(10**6)
    status = await _upload(monkeypatch, model)
    assert status["status"] == "failed"
    assert "cut off" in status["error"], status["error"]
