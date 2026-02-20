import json

import pytest

from src.llm.config import TaskProviderMapping
from src.llm.failover import FailoverManager
from src.llm.health import HealthMonitor
from src.llm.interface import AnalysisContext, LLMConfig, TaskType
from src.llm.providers.mock import MockProvider
from src.services.resume_service import ResumeService, RESUME_EXTRACTION_SYSTEM_PROMPT


class FakeStorage:
    """Fake storage service for testing."""

    def __init__(self):
        self.files: dict[str, bytes] = {}

    def upload_file(self, file_content: bytes, content_type: str, original_filename: str) -> str:
        import uuid
        ref = str(uuid.uuid4())
        self.files[ref] = file_content
        return ref


def _mock_failover_manager() -> FailoverManager:
    """Create a FailoverManager wired to use MockProvider for all task types."""
    health = HealthMonitor()
    fm = FailoverManager(health)
    # Override all task mappings to use mock provider
    for task_type in TaskType:
        fm.update_mapping(TaskProviderMapping(
            task_type=task_type,
            primary_provider="mock",
            primary_model="mock-v1",
            fallback_provider="mock",
            fallback_model="mock-v1",
        ))
    return fm


class TestResumeService:
    @pytest.fixture
    def service(self) -> ResumeService:
        storage = FakeStorage()
        fm = _mock_failover_manager()
        return ResumeService(storage=storage, llm=fm)

    def test_extract_text_from_pdf(self, service: ResumeService):
        import fitz

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Jane Smith\nSenior Engineer at FAANG")
        pdf_bytes = doc.tobytes()
        doc.close()

        text = service.extract_text(pdf_bytes)
        assert "Jane Smith" in text
        assert "Senior Engineer" in text

    def test_extract_text_from_docx(self, service: ResumeService):
        from docx import Document
        from io import BytesIO

        doc = Document()
        doc.add_paragraph("Bob Johnson")
        doc.add_paragraph("Tech Lead")
        buf = BytesIO()
        doc.save(buf)

        text = service.extract_text(buf.getvalue())
        assert "Bob Johnson" in text
        assert "Tech Lead" in text

    def test_extract_text_unknown_format_raises(self, service: ResumeService):
        with pytest.raises(ValueError, match="Unsupported file type"):
            service.extract_text(b"random binary content that is not pdf or docx")

    @pytest.mark.asyncio
    async def test_extract_structured_data_with_mock(self, service: ResumeService):
        """Test that mock LLM returns structured resume data."""
        result = await service.extract_structured_data("John Doe\nSoftware Engineer\n5 years experience")
        assert "name" in result
        assert "tech_stack" in result
        assert "verifiable_claims" in result

    @pytest.mark.asyncio
    async def test_process_resume_pdf_pipeline(self, service: ResumeService):
        """Test full pipeline: PDF → text → structured extraction."""
        import fitz

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Alice Developer\nStaff Engineer\nBuilt systems at scale")
        pdf_bytes = doc.tobytes()
        doc.close()

        result = await service.process_resume(pdf_bytes)
        assert "_raw_text" in result
        assert "Alice Developer" in result["_raw_text"]
        # Mock provider returns structured data
        assert "name" in result

    @pytest.mark.asyncio
    async def test_process_resume_empty_text_raises(self, service: ResumeService):
        """Empty PDF should raise an error."""
        import fitz

        doc = fitz.open()
        doc.new_page()  # empty page, no text
        pdf_bytes = doc.tobytes()
        doc.close()

        with pytest.raises(ValueError, match="No text could be extracted"):
            await service.process_resume(pdf_bytes)


class TestMockProviderResponses:
    """Verify the mock provider returns correct canned responses for resume extraction."""

    @pytest.mark.asyncio
    async def test_mock_resume_extraction(self):
        provider = MockProvider()
        context = AnalysisContext(
            system_prompt="test",
            user_content="test resume",
            task_type=TaskType.RESUME_EXTRACTION,
        )
        config = LLMConfig(response_format="json")
        result = await provider.analyze(context, config)

        assert result.parsed is not None
        assert result.parsed["name"] == "Test Candidate"
        assert "Python" in result.parsed["tech_stack"]
        assert len(result.parsed["verifiable_claims"]) > 0
